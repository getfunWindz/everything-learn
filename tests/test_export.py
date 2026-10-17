import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
import pytest
from export import md_to_html, export_docx

SAMPLE_MD = """# 标题一

> 引用内容

| 概念 | 解释 |
|------|------|
| A | 甲 |

- 列表项1
- 列表项2

**加粗** 和 `行内代码`
"""

def test_md_to_html_basic():
    html = md_to_html(SAMPLE_MD)
    assert "<h1>标题一</h1>" in html
    assert "<blockquote>引用内容</blockquote>" in html
    assert "<table>" in html and "甲" in html
    assert "<ul>" in html and "列表项1" in html
    assert "<strong>加粗</strong>" in html
    assert "<code>行内代码</code>" in html

def test_md_to_html_escape():
    html = md_to_html("# <script>alert(1)</script>")
    assert "<script>" not in html  # XSS 转义
    assert "&lt;script&gt;" in html

def test_export_docx(tmp_path):
    pytest.importorskip("docx")
    out = tmp_path / "r.docx"
    export_docx(SAMPLE_MD, str(out))
    assert out.exists()
    from docx import Document
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "标题一" in texts
    assert "甲" in " ".join(texts) or any("甲" in c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_md_to_txt_strips_markdown():
    from export import md_to_txt
    t = md_to_txt(SAMPLE_MD)
    assert "标题一" in t and "#" not in t.split("\n")[0]
    assert "加粗" in t and "**" not in t and "`" not in t
    assert "A\t甲" in t or "A | 甲" in t


def test_export_report_formats(tmp_path):
    from export import export_report
    d = tmp_path / "item"
    d.mkdir()
    (d / "测试学习报告.md").write_text(SAMPLE_MD, encoding="utf-8")
    p_txt = export_report(str(d), "txt")
    assert p_txt.endswith(".txt") and "标题一" in open(p_txt, encoding="utf-8").read()
    p_doc = export_report(str(d), "doc")
    head = open(p_doc, encoding="utf-8").read(200)
    assert p_doc.endswith(".doc") and "<html" in head.lower()
    p_html = export_report(str(d), "html")
    assert p_html.endswith(".html") and "<h1>" in open(p_html, encoding="utf-8").read()
